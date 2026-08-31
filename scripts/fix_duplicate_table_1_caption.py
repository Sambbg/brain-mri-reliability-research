"""
Remove the stale second caption under Table 1.

Reads:  reports/Gonzalves_BrainMRI_Reliability_Paper_final.docx
Writes: reports/Gonzalves_BrainMRI_Reliability_Paper_final.docx  (in place)

The related-works table ended up with two captions. insert_table_1.py added
"Table 1. Reliability dimensions reported by representative..." to describe the
table it built, but the original placeholder caption, "Table 1. Summary of
related works...", was left in place beneath it. The document therefore had two
paragraphs numbered Table 1 against a single table, and one more caption than
it has tables.

The kept caption is the one that describes the table actually present: its
columns are study, task and dataset, best internal result, split, and the
reliability dimensions, which is what the retained caption names. The removed
caption describes a columns layout the table does not have.

Idempotent: does nothing if the stale caption is absent.
"""

from pathlib import Path
import re
import sys

from docx import Document

DOCX = Path("reports/Gonzalves_BrainMRI_Reliability_Paper_final.docx")

KEEP_PREFIX = "Table 1. Reliability dimensions"
REMOVE_PREFIX = "Table 1. Summary of related works"


def main():
    if not DOCX.exists():
        sys.exit(f"Manuscript not found: {DOCX}")

    document = Document(str(DOCX))

    kept = [p for p in document.paragraphs if p.text.strip().startswith(KEEP_PREFIX)]
    stale = [p for p in document.paragraphs if p.text.strip().startswith(REMOVE_PREFIX)]

    if not stale:
        print("No stale 'Summary of related works' caption found; nothing to do.")
        return

    if not kept:
        sys.exit(
            "Refusing to delete the only Table 1 caption: the retained caption "
            f"({KEEP_PREFIX!r}) is not present. Check the document by hand."
        )

    for para in stale:
        print(f"Removing: {para.text.strip()[:90]}")
        para._element.getparent().remove(para._element)

    document.save(str(DOCX))

    captions = [
        p.text.strip() for p in Document(str(DOCX)).paragraphs
        if re.match(r"^Table \d+\.", p.text.strip())
    ]
    numbers = [int(re.match(r"^Table (\d+)\.", c).group(1)) for c in captions]
    duplicates = sorted({n for n in numbers if numbers.count(n) > 1})

    print(f"\nRemoved {len(stale)} caption(s).")
    print(f"Captions now: {len(captions)}; tables: {len(Document(str(DOCX)).tables)}")
    print(f"Duplicate caption numbers remaining: {duplicates or 'none'}")


if __name__ == "__main__":
    main()
