"""
Set the funding statement in the Acknowledgments and the Declarations.

Reads:  reports/Gonzalves_BrainMRI_Reliability_Paper_final.docx
Writes: reports/Gonzalves_BrainMRI_Reliability_Paper_final.docx  (in place)

Two different operations, because the two sections were not in the same state.

The Acknowledgments carried a bracketed placeholder mid-paragraph, which is
replaced in place.

The Declarations had no funding entry at all -- not a placeholder, an absence.
A `Funding.` entry is added, matching the bold-label format of the surrounding
entries and placed before `Competing interests.`, which is the conventional
order. This adds content rather than replacing a placeholder; it is done because
a Declarations block that omits funding is incomplete for submission, but it is
the one change here that is an addition and it is called out in the run output.

Idempotent: skips whichever part is already done.
"""

from pathlib import Path
import sys

from docx import Document
from docx.shared import Pt

DOCX = Path("reports/Gonzalves_BrainMRI_Reliability_Paper_final.docx")

FUNDING_TEXT = "The authors received no specific funding for this work."

PLACEHOLDER = (
    "[State the funding source and cost centre number here, or: "
    "The authors received no specific funding for this work.]"
)

ACK_PREFIX = "The authors express their gratitude"
DECLARATIONS_ANCHOR = "Competing interests."
FUNDING_LABEL = "Funding. "


def replace_placeholder(document):
    target = next(
        (p for p in document.paragraphs if p.text.strip().startswith(ACK_PREFIX)),
        None,
    )
    if target is None:
        sys.exit(f"Could not find the Acknowledgments paragraph ({ACK_PREFIX!r}).")

    if PLACEHOLDER not in target.text:
        print("Acknowledgments: placeholder already replaced, skipping.")
        return False

    # One run holds the whole paragraph, so a substring swap preserves formatting.
    for run in target.runs:
        if PLACEHOLDER in run.text:
            run.text = run.text.replace(PLACEHOLDER, FUNDING_TEXT)
            print("Acknowledgments: placeholder replaced.")
            return True

    sys.exit("The placeholder spans multiple runs; replace it by hand.")


def add_funding_declaration(document):
    if any(p.text.strip().startswith(FUNDING_LABEL.strip()) for p in document.paragraphs):
        print("Declarations: a Funding entry is already present, skipping.")
        return False

    anchor = next(
        (p for p in document.paragraphs if p.text.strip().startswith(DECLARATIONS_ANCHOR)),
        None,
    )
    if anchor is None:
        sys.exit(f"Could not find the {DECLARATIONS_ANCHOR!r} declaration.")

    label_run, body_run = anchor.runs[0], anchor.runs[1]

    new_element = anchor._p.makeelement(anchor._p.tag, {})
    anchor._p.addprevious(new_element)

    from docx.text.paragraph import Paragraph
    para = Paragraph(new_element, anchor._parent)
    para.alignment = anchor.alignment
    para.paragraph_format.space_after = anchor.paragraph_format.space_after

    label = para.add_run(FUNDING_LABEL)
    label.bold = True
    label.font.name = label_run.font.name
    label.font.size = label_run.font.size

    body = para.add_run(FUNDING_TEXT)
    body.font.name = body_run.font.name
    body.font.size = body_run.font.size

    print("Declarations: added a Funding entry (this section had none).")
    return True


def main():
    if not DOCX.exists():
        sys.exit(f"Manuscript not found: {DOCX}")

    document = Document(str(DOCX))

    changed = replace_placeholder(document)
    changed = add_funding_declaration(document) or changed

    if not changed:
        print("\nNothing to do.")
        return

    document.save(str(DOCX))
    print(f"\nSaved: {DOCX}")


if __name__ == "__main__":
    main()
