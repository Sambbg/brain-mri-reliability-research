"""
Add the leakage-free sensitivity analysis to sections 2.2 and 3.2.

Reads:  reports/Gonzalves_BrainMRI_Reliability_Paper_final.docx
Writes: reports/Gonzalves_BrainMRI_Reliability_Paper_final.docx  (in place)

Section 2.2 already carries the measured patient-level residual. This adds what
that residual does to the reported numbers, and why the split was not rebuilt.
Section 3.2 gets the clean-subset figures beside the headline ones and the
memorisation gap that shows the contamination is real.

Figures are from reports/experiments/consolidated/tables/leakage_free_subset.md.

Two drafting constraints, both deliberate.

The signal-to-noise decline on the subset (2.52 to 2.09 for macro-F1) is not
mentioned. Most of it is mechanical: mean seed SD rises by a factor of 1.304
while the partition shrinks by 1.285, because the excluded images are ones every
seed classified correctly, so removing them takes away almost no seed-to-seed
disagreement while shrinking the denominator. Quoting the decline without that
explanation would read as increased seed sensitivity, which it is not. It stays
in the supplementary report where the explanation sits with it.

The memorisation gap is included. Every architecture classifies the excluded
images 2.4 to 3.1 points more accurately than the rest, in every run. That is
the evidence distinguishing genuine contamination from a perceptual-hashing
artefact, and it is the strongest single line in the analysis.

Idempotent: aborts if either paragraph is already present.
"""

from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

DOCX = Path("reports/Gonzalves_BrainMRI_Reliability_Paper_final.docx")

METHODS_ANCHOR_PREFIX = "Because D1 is distributed without patient-level identifiers"
RESULTS_ANCHOR_PREFIX = "This is the first substantive result of the study"

METHODS_SENTINEL = "was measured rather than assumed"
RESULTS_SENTINEL = "no same-patient counterpart in training, the three architectures"

METHODS_PARAGRAPH = (
    "The effect of that residual on the reported figures was measured rather "
    "than assumed. Recomputing every internal metric for all fifteen runs on the "
    "818 test images that have no same-patient counterpart in training lowers "
    "mean test macro-F1 by 0.0039 for ResNet18, 0.0056 for EfficientNet-B0 and "
    "0.0059 for ViT-B/16, and lowers accuracy by between 0.0052 and 0.0068. The "
    "shift is close to uniform across the three architectures and comparable in "
    "magnitude to the mean within-architecture variation across seeds of 0.0046, "
    "so it does not alter the comparisons drawn between them. The split was "
    "accordingly not rebuilt at the looser threshold. Regenerating it would "
    "invalidate all fifteen trained checkpoints and every result derived from "
    "them, which a correction of this size, measured against that much "
    "seed-induced variation, does not warrant. Uncorrected figures are reported "
    "throughout, with the leakage-free values given alongside them in Section 3.2."
)

RESULTS_PARAGRAPH = (
    "This conclusion is not an artefact of the residual patient-level "
    "contamination quantified in Section 2.2. Recomputed on the 818 test images "
    "with no same-patient counterpart in training, the three architectures give "
    "mean test macro-F1 of 0.9623 ± 0.0084, 0.9617 ± 0.0051 and 0.9497 "
    "± 0.0045. The ranking remains unstable in the same way and to the same "
    "degree: the same two orderings occur across the five seeds, produced by the "
    "same seeds, and the margin of 0.0011 by which EfficientNet-B0 leads ResNet18 "
    "on the full partition becomes a deficit of 0.0005 on the subset. That the "
    "contamination is genuine rather than an artefact of perceptual hashing is "
    "evident in the excluded images themselves, which every architecture "
    "classifies between 2.4 and 3.1 percentage points more accurately than the "
    "remainder of the partition, in every one of the fifteen runs."
)


def insert_after(anchor, text):
    """New body paragraph immediately after the anchor, matching its formatting."""
    template = anchor.runs[0]
    font_name = template.font.name
    font_size = template.font.size

    new_element = anchor._p.makeelement(anchor._p.tag, {})
    anchor._p.addnext(new_element)

    from docx.text.paragraph import Paragraph
    para = Paragraph(new_element, anchor._parent)
    para.alignment = anchor.alignment or WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.space_after = anchor.paragraph_format.space_after or Pt(8)

    run = para.add_run(text)
    run.font.name = font_name
    run.font.size = font_size if font_size is not None else Pt(11)

    return para


def main():
    if not DOCX.exists():
        sys.exit(f"Manuscript not found: {DOCX}")

    document = Document(str(DOCX))
    text = "\n".join(p.text for p in document.paragraphs)

    if METHODS_SENTINEL in text or RESULTS_SENTINEL in text:
        sys.exit("The sensitivity paragraphs are already present; nothing to do.")

    inserted = []

    for prefix, paragraph_text, label in (
        (METHODS_ANCHOR_PREFIX, METHODS_PARAGRAPH, "2.2"),
        (RESULTS_ANCHOR_PREFIX, RESULTS_PARAGRAPH, "3.2"),
    ):
        anchor = next(
            (p for p in document.paragraphs if p.text.strip().startswith(prefix)),
            None,
        )
        if anchor is None:
            sys.exit(f"Could not find the section {label} anchor: {prefix!r}")

        insert_after(anchor, paragraph_text)
        inserted.append((label, prefix))

    document.save(str(DOCX))

    for label, prefix in inserted:
        print(f"Section {label}: inserted after {prefix[:60]!r}...")
    print(f"\nSaved: {DOCX}")


if __name__ == "__main__":
    main()
