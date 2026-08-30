"""
Insert the seven paper figures into the manuscript and upgrade their captions.

Replaces each boxed grey "INSERT FIGURE N" placeholder paragraph with the
corresponding PNG, centred and sized to the width specified below, and replaces
the short italic caption directly beneath it with the fuller caption from
figure_captions.md.

Only the width is set on each image, so Word scales the height and the aspect
ratio is preserved. The placeholder paragraph is replaced rather than emptied:
the grey box is a paragraph border (w:pBdr) on the placeholder itself, so
reusing that paragraph would leave the image sitting inside the box.

The six "INSERT TABLE N" placeholders are matched by a different prefix and are
deliberately untouched, as are all body text, headings, equations and references.

Writes to a new file; the input document is never modified.
"""

from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

DOCX_IN = Path("reports/Gonzalves_BrainMRI_Reliability_Paper.docx")
DOCX_OUT = Path("reports/Gonzalves_BrainMRI_Reliability_Paper_with_figures.docx")

FIGURE_DIR = Path("reports/experiments/consolidated/paper_figures")
CAPTIONS_MD = FIGURE_DIR / "figure_captions.md"

# figure number -> (file name, width in inches)
FIGURES = {
    1: ("figure_1_system_architecture.png", 6.5),
    2: ("figure_2_internal_seed_spread.png", 4.0),
    3: ("figure_3_reliability_diagrams.png", 6.5),
    4: ("figure_4_internal_vs_shifted_domain.png", 5.5),
    5: ("figure_5_prediction_distribution.png", 6.5),
    6: ("figure_6_confidence_inversion.png", 4.5),
    7: ("figure_7_correlation_decomposition.png", 6.5),
}

CAPTION_POINT_SIZE = 10

# Matches the placeholder spacing so the figures sit in the flow like the
# placeholders did. 160 twips = 8pt.
IMAGE_SPACING_TWIPS = 160

PLACEHOLDER_RE = re.compile(r"^INSERT\s+FIGURE\s+(\d+)\b", re.IGNORECASE)


def join_wrapped_lines(block):
    """
    Rejoin a hard-wrapped markdown block into one paragraph.

    A line ending in a hyphen followed by a lowercase continuation is a split
    compound word ("between-" / "architecture"), so it joins without a space.
    Everything else joins with one space.
    """
    lines = [line.strip() for line in block.split("\n") if line.strip()]

    if not lines:
        return ""

    joined = lines[0]

    for line in lines[1:]:
        if joined.endswith("-") and line[:1].islower():
            joined += line
        else:
            joined += " " + line

    return joined


def load_captions(path):
    """Parse figure_captions.md into {figure number: caption text}."""
    if not path.exists():
        raise FileNotFoundError(f"Captions file not found: {path}")

    text = path.read_text(encoding="utf-8")
    captions = {}

    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        match = re.match(r"\*\*Figure (\d+)\.\*\*", block)

        if match:
            captions[int(match.group(1))] = join_wrapped_lines(block)

    return captions


def split_bold_segments(text):
    """
    Split markdown text into (segment, is_bold) pairs on ** ** markers.

    The figure label and, in figure 2, the axis-truncation warning are emphasised
    in the source. Flattening them would lose emphasis the author put there
    deliberately -- the truncation note is exactly the thing reviewers query.
    """
    segments = []

    for i, part in enumerate(text.split("**")):
        if part:
            segments.append((part, i % 2 == 1))

    return segments


def clear_paragraph_runs(paragraph):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def set_caption(paragraph, text):
    """Rewrite a caption paragraph as italic 10pt, left-aligned."""
    clear_paragraph_runs(paragraph)

    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = None
    paragraph.paragraph_format.first_line_indent = None

    for segment, is_bold in split_bold_segments(text):
        run = paragraph.add_run(segment)
        run.italic = True
        run.bold = is_bold
        run.font.size = Pt(CAPTION_POINT_SIZE)


def insert_image(placeholder, image_path, width_inches):
    """
    Put a centred image where the placeholder paragraph was.

    A fresh paragraph is inserted before the placeholder and the placeholder is
    then removed, so none of its border or shading survives.
    """
    image_paragraph = placeholder.insert_paragraph_before()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(IMAGE_SPACING_TWIPS / 20)
    image_paragraph.paragraph_format.space_after = Pt(IMAGE_SPACING_TWIPS / 20)

    run = image_paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))

    placeholder._element.getparent().remove(placeholder._element)


def main():
    if not DOCX_IN.exists():
        raise FileNotFoundError(f"Manuscript not found: {DOCX_IN}")

    captions = load_captions(CAPTIONS_MD)

    missing_images = [
        name for name, _ in FIGURES.values()
        if not (FIGURE_DIR / name).exists()
    ]
    if missing_images:
        raise FileNotFoundError(f"Figure files not found: {missing_images}")

    missing_captions = sorted(set(FIGURES) - set(captions))
    if missing_captions:
        raise ValueError(
            f"figure_captions.md has no caption for figures: {missing_captions}"
        )

    document = Document(str(DOCX_IN))

    # Resolve every target before mutating anything. Paragraph objects hold XML
    # element references that stay valid across inserts and deletes elsewhere in
    # the document, but positional indices would not.
    targets = {}
    duplicates = []
    paragraphs = document.paragraphs

    for index, paragraph in enumerate(paragraphs):
        match = PLACEHOLDER_RE.match(paragraph.text.strip())

        if not match:
            continue

        number = int(match.group(1))

        if number in targets:
            duplicates.append(number)
            continue

        caption_paragraph = paragraphs[index + 1] if index + 1 < len(paragraphs) else None

        targets[number] = {
            "placeholder": paragraph,
            "caption": caption_paragraph,
            "caption_before": caption_paragraph.text.strip() if caption_paragraph else None,
            "index": index,
        }

    replaced = []
    not_found = []
    caption_warnings = []

    for number in sorted(FIGURES):
        name, width = FIGURES[number]

        if number not in targets:
            not_found.append(number)
            continue

        target = targets[number]
        caption_paragraph = target["caption"]

        if caption_paragraph is None or not caption_paragraph.text.strip().startswith(
            f"Figure {number}."
        ):
            caption_warnings.append(number)
            caption_paragraph = None

        insert_image(target["placeholder"], FIGURE_DIR / name, width)

        if caption_paragraph is not None:
            set_caption(caption_paragraph, captions[number])

        replaced.append({
            "number": number,
            "file": name,
            "width_inches": width,
            "paragraph_index": target["index"],
            "caption_before": target["caption_before"],
            "caption_after": captions[number],
        })

    remaining = [
        p.text.strip() for p in document.paragraphs
        if p.text.strip().upper().startswith("INSERT ")
    ]

    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(DOCX_OUT))

    print(f"Input:  {DOCX_IN}")
    print(f"Output: {DOCX_OUT}\n")

    print(f"Replaced {len(replaced)} of {len(FIGURES)} figure placeholders:\n")
    for entry in replaced:
        print(f"  FIGURE {entry['number']} -> {entry['file']} @ {entry['width_inches']} in")
        print(f"    caption before: {entry['caption_before'][:80]}...")
        print(f"    caption after:  {entry['caption_after'][:80]}...")

    if not_found:
        print(f"\nNOT FOUND: figure placeholders {not_found}")

    if duplicates:
        print(f"\nWARNING duplicate placeholders ignored: {duplicates}")

    if caption_warnings:
        print(
            f"\nWARNING no matching caption paragraph beneath figures "
            f"{caption_warnings}; image inserted, caption left alone"
        )

    print(f"\nPlaceholders remaining in the document ({len(remaining)}):")
    for text in remaining:
        print(f"  {text[:90]}")

    if not_found or caption_warnings:
        sys.exit(1)


if __name__ == "__main__":
    main()
