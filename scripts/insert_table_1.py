#!/usr/bin/env python3
"""
Insert Table 1 (literature positioning) into the manuscript.

Unlike Tables 2-6, this is not built from experimental artefacts. Every cell was
read from the full text of the cited paper, and the evidence for each is recorded
in reports/experiments/TABLE_1_EVIDENCE.md.

Four corrections to figures commonly quoted from these papers are applied here
and explained in the table footnote:

  - Disci et al. report 98.73%, which has support = 7023 (train + test) and is a
    sample-weighted average of training and testing accuracy. The test accuracy
    is 95.27%.
  - Shah et al. report 98.87% on a binary tumour/non-tumour task, on the
    validation split, not a multi-class test set.
  - Reyes and Sanchez ran five random seeds but report only the maximum.
  - Vimala et al. is omitted: its reported figures are numerically identical to
    Zulfiqar et al. on the same dataset and pipeline, and the relationship
    between the two papers is unresolved.

Run from the repo root:
    python scripts/insert_table_1.py

Reads:   reports/Gonzalves_BrainMRI_Reliability_Paper_final.docx
Writes:  reports/Gonzalves_BrainMRI_Reliability_Paper_final.docx  (in place)
"""

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

DOC = Path("reports/Gonzalves_BrainMRI_Reliability_Paper_final.docx")

FONT = "Times New Roman"
BODY_PT = 8
HEAD_PT = 8

HEADERS = [
    "Study",
    "Task and dataset",
    "Best internal result",
    "Split",
    "Calibration",
    "Leakage or overlap audit",
    "External or shifted-domain validation",
    "Repeated training runs",
]

ROWS = [
    [
        "Shah et al. (2022) [6]",
        "Binary tumour / non-tumour; Kaggle collection described as a BraTS-derived subset",
        "98.87% validation accuracy \u1d43",
        "Custom 80/20 plus 60 held-out images",
        "No",
        "No",
        "No",
        "No",
    ],
    [
        "Zulfiqar et al. (2023) [7]",
        "Three-class; Figshare CE-MRI (3,064 slices, 233 patients)",
        "98.86% test accuracy; mean F1 98.71%",
        "Custom 80/20, slice level",
        "No",
        "No \u1d47",
        "Yes, but the target set shares an upstream source \u1d9c",
        "No",
    ],
    [
        "Reyes and S\u00e1nchez (2024) [30]",
        "Three-class and four-class; Figshare CE-MRI and a Kaggle collection",
        "98.7% test accuracy",
        "Custom 80/10/10 on each dataset",
        "No",
        "Partly: cross-dataset overlap quantified; no patient-level split \u1d48",
        "No: each dataset trained and tested within itself",
        "Yes, five seeds, but only the maximum reported \u1d49",
    ],
    [
        "Disci et al. (2025) [5]",
        "Four-class; aggregated Kaggle collection (7,023 images)",
        "95.27% test accuracy; macro-F1 0.9491 \u1da0",
        "Distributed split as published",
        "No",
        "No",
        "No, acknowledged as a limitation",
        "No, and no cross-validation",
    ],
    [
        "Elhadidy et al. (2025) [8]",
        "Four-class; Kaggle collection, augmented to 9,749 images",
        "98.72% test accuracy",
        "Custom 80/20 with 5-fold cross-validation \u1d4d",
        "No",
        "No",
        "No",
        "No; cross-validation only",
    ],
    [
        "This study",
        "Four-class internal; two glioma-focused shifted-domain probes",
        "0.9673 \u00b1 0.0043 test macro-F1 (best architecture mean)",
        "Leakage-aware, near-duplicate groups assigned whole",
        "Yes: ECE, NLL, Brier, confidence gap, before and after temperature scaling",
        "Yes: exact and perceptual audit of every candidate; one rejected",
        "Shifted-domain only, reported as behaviour rather than accuracy",
        "Yes: five seeds per architecture, all values reported as mean \u00b1 SD",
    ],
]

FOOTNOTE = (
    "Table 1. Reliability dimensions reported by representative high-performing "
    "brain MRI tumour classification studies. Every cell was determined from the "
    "full text of the cited paper; the supporting evidence is recorded in the "
    "supplementary material. Internal performance is reported by every study, "
    "whereas calibration analysis, leakage or overlap auditing, independent "
    "external validation and repeated training runs are almost entirely absent. "
    "\u1d43 Binary task on the validation split, not a multi-class test set, and "
    "therefore not directly comparable with the other rows. "
    "\u1d47 Patient identifiers are present in the source collection and were not "
    "used; slices from one patient may fall on both sides of the split. "
    "\u1d9c Validation was performed on a Kaggle collection subsequently documented "
    "as sharing more than 2,260 images with the training source [30]; the study "
    "makes no independence claim or check. "
    "\u1d48 The overlap is quantified but the detection method is not described. "
    "\u1d49 Five random initialisations were run and the highest accuracy reported, "
    "which biases the headline figure upward relative to a single run. "
    "\u1da0 The value of 98.73% reported in the source is a sample-weighted average "
    "over training and test partitions (support 7,023) rather than a test metric; "
    "the test accuracy is given here. "
    "\u1d4d Augmentation appears to precede the split, so augmented copies of one "
    "source image may fall on both sides. "
    "A sixth study meeting the inclusion criteria was omitted: its reported "
    "figures are numerically identical to those of Zulfiqar et al. on the same "
    "dataset and pipeline, and the relationship between the two reports is "
    "unresolved."
)


def apply_borders(table, size=4, colour="666666"):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), colour)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def shade_cell(cell, fill="E8E8E8"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def style_cell(cell, text, bold=False, align=None, size=BODY_PT):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.space_before = Pt(0)
    run = para.add_run(str(text))
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold


def main():
    if not DOC.exists():
        sys.exit(f"Manuscript not found: {DOC}\n"
                 "Run insert_paper_tables.py first.")

    doc = Document(str(DOC))

    anchor = next((p for p in doc.paragraphs
                   if re.match(r"\s*INSERT TABLE 1\b", p.text)), None)
    if anchor is None:
        sys.exit("No 'INSERT TABLE 1' placeholder found. "
                 "It may already have been replaced.")

    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    apply_borders(table)

    for i, head in enumerate(HEADERS):
        style_cell(table.rows[0].cells[i], head, bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER, size=HEAD_PT)
        shade_cell(table.rows[0].cells[i])

    for row_vals in ROWS:
        cells = table.add_row().cells
        last = row_vals[0] == "This study"
        for i, val in enumerate(row_vals):
            style_cell(cells[i], val, bold=last)
            if last:
                shade_cell(cells[i], "F2F2F2")

    widths = [13, 17, 14, 13, 12, 14, 15, 12]
    total = sum(widths)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(6.5 * widths[i] / total)

    anchor._p.addnext(table._tbl)

    # replace the placeholder paragraph with the caption
    anchor.text = ""
    anchor.alignment = WD_ALIGN_PARAGRAPH.LEFT
    anchor.paragraph_format.space_before = Pt(6)
    for attr in ("top", "bottom", "left", "right"):
        pass  # border removal below
    pPr = anchor._p.get_or_add_pPr()
    for bdr in pPr.findall(qn("w:pBdr")):
        pPr.remove(bdr)
    run = anchor.add_run(FOOTNOTE)
    run.font.name = FONT
    run.font.size = Pt(8)
    run.italic = True

    # the caption belongs below the table, so move the paragraph after it
    table._tbl.addnext(anchor._p)

    doc.save(str(DOC))
    print(f"Table 1 inserted: {len(ROWS)} rows x {len(HEADERS)} columns")
    print(f"Saved: {DOC}")
    print("\nNote: the existing short caption paragraph beneath the old")
    print("placeholder may now be redundant. Check and delete if so.")


if __name__ == "__main__":
    main()
